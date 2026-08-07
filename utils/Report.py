import os
import html
import json
from datetime import datetime, date
from decimal import Decimal

from dto.StandardYAMLAnalysisTemplate import StandardFlowResult, StandardStepResult

"""
测试结果报告生成器 — 轻量 HTML + Excel(.xlsx)

用法:
    from utils.Report import generate
    generate(result, "out")                     # 写入 out/report.html + out/report.xlsx
    generate(result)                            # 缺省写入 reports/<时间戳>/
"""

# ═══════════════════════════════════════════════════════════════
#  公共入口
# ═══════════════════════════════════════════════════════════════
def generate(result: StandardFlowResult, out_dir: str | None = None) -> str:
    """
    生成 HTML + Excel 报告, 返回输出目录路径。
    out_dir 缺省 → reports/<YYYYMMDD_HHMMSS>/
    """
    if out_dir is None:
        out_dir = os.path.join("reports", datetime.now().strftime("%Y%m%d_%H%M%S"))
    os.makedirs(out_dir, exist_ok=True)
    generate_html(result, os.path.join(out_dir, "report.html"))
    generate_excel(result, os.path.join(out_dir, "report.xlsx"))
    return out_dir

# ═══════════════════════════════════════════════════════════════
#  序列化 / 截断辅助
# ═══════════════════════════════════════════════════════════════
def _jsonable(obj):
    """递归转换为 JSON 可序列化对象 (Decimal/datetime/bytes 等)。"""
    if isinstance(obj, dict):
        return {str(k): _jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_jsonable(v) for v in obj]
    if isinstance(obj, bool): return obj
    if isinstance(obj, (int, float)): return obj
    if isinstance(obj, Decimal): return float(obj)
    if isinstance(obj, (datetime, date)): return obj.isoformat()
    if isinstance(obj, bytes): return obj.decode("utf-8", errors="replace")
    if obj is None: return None
    return str(obj)

_EXCEL_MAX_CELL: int = 32767   # Excel 单单元格最大字符数(平台硬限制)

def _truncate(s: str, n: int | None = None) -> str:
    """超长截断, 附省略标记。n=None → 不截断, 返回原串。"""
    if s is None: return ""
    if n is None or len(s) <= n: return s
    marker: str = f" ...(截断 {len(s)} 字符)"
    cut: int = max(0, n - len(marker))
    return s[:cut] + f" ...(截断 {len(s) - cut} 字符)"

def _json_pretty(obj, maxlen: int | None = None) -> str:
    """对象 → 缩进 JSON 字符串。maxlen=None → 不截断。"""
    return _truncate(json.dumps(_jsonable(obj), ensure_ascii=False, indent=2), maxlen)

def _esc(s) -> str:
    """HTML 转义。"""
    return html.escape(str(s) if s is not None else "")

# ═══════════════════════════════════════════════════════════════
#  HTML 报告
# ═══════════════════════════════════════════════════════════════
_CSS = """
body { font-family:'Segoe UI','Microsoft YaHei',sans-serif; margin:24px; color:#1f2937; background:#f3f4f6; overflow-x:hidden; }
h1 { font-size:22px; margin:0 0 4px; overflow-wrap:anywhere; word-break:break-all; }
h2 { font-size:17px; margin:28px 0 10px; border-bottom:1px solid #d1d5db; padding-bottom:6px; overflow-wrap:anywhere; word-break:break-all; }
.meta { color:#6b7280; font-size:13px; margin-bottom:16px; overflow-wrap:anywhere; word-break:break-all; }
.cards { display:flex; gap:14px; flex-wrap:wrap; margin:16px 0 8px; }
.card { background:#fff; border:1px solid #e5e7eb; border-radius:10px; padding:14px 20px; min-width:110px; text-align:center; box-shadow:0 1px 2px rgba(0,0,0,.04); }
.card .num { font-size:26px; font-weight:700; }
.card .lbl { font-size:12px; color:#6b7280; margin-top:2px; }
.card.ok .num { color:#16a34a; }
.card.bad .num { color:#dc2626; }
.step { background:#fff; border:1px solid #e5e7eb; border-radius:10px; padding:14px 18px; margin:12px 0; box-shadow:0 1px 2px rgba(0,0,0,.04); max-width:100%; min-width:0; overflow:hidden; }
.step .head { display:flex; align-items:center; gap:10px; flex-wrap:wrap; }
.badge { font-size:12px; font-weight:700; padding:3px 10px; border-radius:999px; color:#fff; flex:none; }
.badge.pass { background:#16a34a; }
.badge.fail { background:#dc2626; }
.step .name { font-size:15px; font-weight:600; min-width:0; overflow-wrap:anywhere; word-break:break-all; }
.err { color:#dc2626; font-size:13px; margin:8px 0 0; overflow-wrap:anywhere; word-break:break-all; }
.err::before { content:'⚠ '; }
table { table-layout:fixed; border-collapse:collapse; width:100%; margin:8px 0; font-size:13px; }
th,td { border:1px solid #e5e7eb; padding:6px 10px; text-align:left; vertical-align:top; overflow-wrap:anywhere; word-break:break-all; }
th { background:#f9fafb; font-weight:600; }
pre { background:#f8fafc; border:1px solid #e5e7eb; border-radius:6px; padding:10px; font-size:12px; overflow-x:auto; white-space:pre-wrap; word-break:break-all; overflow-wrap:anywhere; max-width:100%; box-sizing:border-box; margin:6px 0 0; }
.op { margin:10px 0 0; padding-top:10px; border-top:1px dashed #e5e7eb; max-width:100%; min-width:0; overflow:hidden; }
.op .ophead { font-size:13px; font-weight:600; color:#374151; overflow-wrap:anywhere; word-break:break-all; }
.op .ophead .s { color:#6b7280; font-weight:400; }
.op .rows { color:#6b7280; font-size:12px; margin-top:6px; overflow-wrap:anywhere; word-break:break-all; }
.kv { font-size:12px; color:#6b7280; margin:6px 0 0; overflow-wrap:anywhere; word-break:break-all; }
/* ── 通过/失败 卡片可点击 ── */
.card.clickable { cursor:pointer; user-select:none; }
.card.clickable:hover { border-color:#93c5fd; box-shadow:0 2px 8px rgba(0,0,0,.10); }
.card .arrow { font-size:10px; color:#9ca3af; margin-left:2px; }
/* ── 下拉面板: 固定定位, 不撑破页面布局 ── */
.drop { position:fixed; top:76px; left:50%; transform:translateX(-50%); width:min(600px, calc(100vw - 32px)); max-height:calc(100vh - 104px); background:#fff; border:1px solid #d1d5db; border-radius:12px; box-shadow:0 12px 32px rgba(0,0,0,.20); z-index:1000; display:flex; flex-direction:column; overflow:hidden; }
.drop-head { display:flex; align-items:center; justify-content:space-between; gap:8px; padding:10px 14px; border-bottom:1px solid #e5e7eb; font-weight:600; font-size:13px; flex:none; }
.drop-head .cnt { color:#6b7280; font-weight:400; font-size:12px; }
.drop-close { border:none; background:#f3f4f6; border-radius:8px; width:26px; height:26px; cursor:pointer; font-size:16px; line-height:1; color:#374151; flex:none; }
.drop-close:hover { background:#e5e7eb; }
.drop-list { overflow-y:auto; flex:1 1 auto; min-height:48px; padding:6px; }
.drop-item { padding:8px 12px; font-size:12px; border-radius:8px; cursor:pointer; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; margin-bottom:2px; color:#1f2937; }
.drop-item:hover { background:#eef2ff; }
.drop-empty { color:#9ca3af; font-size:13px; padding:14px 10px; text-align:center; }
.target-flash { animation: flashBg 1.6s ease; border-radius:8px; }
@keyframes flashBg { 0%,55% { background:#fef3c7; } 100% { background:transparent; } }
"""

_DROP_JS: str = """
<script>
(function () {
  var CLOSED = 'none';
  function closeAll() {
    document.querySelectorAll('.drop').forEach(function (d) { d.style.display = CLOSED; });
  }
  function toggle(pid) {
    var d = document.getElementById(pid);
    if (!d) return;
    var open = d.style.display !== CLOSED;
    closeAll();
    if (!open) d.style.display = 'flex';
  }
  document.addEventListener('click', function (e) {
    var card = e.target.closest('.card.clickable');
    if (card) { toggle(card.getAttribute('data-drop')); return; }
    if (e.target.closest('.drop')) return;
    closeAll();
  });
  document.addEventListener('keydown', function (e) {
    if (e.key === 'Escape') closeAll();
  });
  document.querySelectorAll('.drop').forEach(function (d) {
    var close = d.querySelector('.drop-close');
    if (close) close.addEventListener('click', function () { d.style.display = CLOSED; });
    d.querySelectorAll('.drop-item').forEach(function (it) {
      it.addEventListener('click', function () {
        var target = document.getElementById(it.getAttribute('data-target'));
        closeAll();
        if (!target) return;
        target.scrollIntoView({ behavior: 'smooth', block: 'start' });
        target.classList.remove('target-flash');
        void target.offsetWidth;
        target.classList.add('target-flash');
        setTimeout(function () { target.classList.remove('target-flash'); }, 1600);
      });
    });
  });
})();
</script>
"""

def generate_html(result: StandardFlowResult, path: str) -> None:
    """渲染自包含 HTML 报告。"""
    total: int = result.total
    passed: int = result.passed
    failed: int = result.failed
    rate: float = round(passed / total * 100, 1) if total else 0.0
    now: str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    parts: list[str] = []
    parts.append("<!DOCTYPE html><html lang='zh'><head><meta charset='utf-8'>")
    parts.append(f"<title>{_esc(result.flow_name)} — 测试报告</title>")
    parts.append(f"<style>{_CSS}</style></head><body>")
    parts.append(f"<h1>测试结果报告</h1>")
    parts.append(f"<div class='meta'>{_esc(result.flow_name)} ｜ 生成时间 {now}</div>")

    parts.append("<div class='cards'>")
    parts.append(f"<div class='card'><div class='num'>{total}</div><div class='lbl'>步骤数</div></div>")
    parts.append(f"<div class='card ok clickable' data-drop='drop-pass' role='button' tabindex='0' title='点击展开通过的操作列表'>"
                 f"<div class='num'>{passed}</div><div class='lbl'>通过 <span class='arrow'>▾</span></div></div>")
    parts.append(f"<div class='card bad clickable' data-drop='drop-fail' role='button' tabindex='0' title='点击展开失败的操作列表'>"
                 f"<div class='num'>{failed}</div><div class='lbl'>失败 <span class='arrow'>▾</span></div></div>")
    parts.append(f"<div class='card'><div class='num'>{rate}%</div><div class='lbl'>通过率</div></div>")
    parts.append(f"<div class='card'><div class='num'>{round(result.duration, 3)}</div><div class='lbl'>耗时(秒)</div></div>")
    parts.append("</div>")

    # 收集 通过/失败 的操作列表(下拉面板用, 与渲染共用同一 op 编号)
    passed_items: list[tuple] = []
    failed_items: list[tuple] = []
    op_index: int = 0
    for i, step in enumerate(result.steps):
        chunk, items, op_index = _render_step_html(step, i, op_index)
        parts.append(chunk)
        passed_items.extend(items["passed"])
        failed_items.extend(items["failed"])

    # 失败但无任何操作 → 补步骤级条目(跳转到步骤卡片)
    for i, step in enumerate(result.steps):
        if not step.operations and not step.passed:
            failed_items.append((f"步骤: {step.name}", step.name, f"step-{i}"))

    parts.append(_drop_panel_html("pass", passed_items, "通过的操作"))
    parts.append(_drop_panel_html("fail", failed_items, "失败的操作"))
    parts.append(_DROP_JS)

    parts.append("</body></html>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))

def _render_step_html(step: StandardStepResult, step_index: int, start_op: int) -> tuple[str, dict, int]:
    """渲染单个步骤卡片。返回 (html, {"passed": 通过项, "failed": 失败项}, 下一 op 编号)。"""
    badge: str = "<span class='badge pass'>PASS</span>" if step.passed else "<span class='badge fail'>FAIL</span>"
    out: list[str] = [f"<div class='step' id='step-{step_index}'><div class='head'>{badge}<span class='name'>{_esc(step.name)}</span></div>"]

    if step.errors:
        out.append("<div>" + "".join(f"<div class='err'>{_esc(e)}</div>" for e in step.errors) + "</div>")

    if step.extracted_vars:
        out.append("<table><colgroup><col style='width:160px'><col></colgroup>"
                   "<tr><th>提取字段</th><th>值</th></tr>")
        for k, v in step.extracted_vars.items():
            out.append(f"<tr><td>{_esc(k)}</td><td><pre>{_esc(_json_pretty(v))}</pre></td></tr>")
        out.append("</table>")

    passed_items: list[tuple] = []
    failed_items: list[tuple] = []
    for op in step.operations:
        op_id: str = f"op-{start_op}"
        out.append(_render_op_html(op, op_id))
        item: tuple = _op_list_item(op, step.name, op_id)
        (passed_items if op.get("passed") else failed_items).append(item)
        start_op += 1

    out.append("</div>")
    return "".join(out), {"passed": passed_items, "failed": failed_items}, start_op

def _render_op_html(op: dict, op_id: str) -> str:
    t: str = op.get("type", "?")
    out: list[str] = [f"<div class='op' id='{op_id}'>"]
    if t == "http":
        out.append(f"<div class='ophead'>HTTP <span class='s'>[{_esc(op.get('method'))}] {_esc(op.get('path'))}"
                   f" → {_esc(op.get('status_code'))} ｜ {_esc(op.get('elapsed'))}s</span></div>")
        if op.get("url"):
            out.append(f"<div class='kv'>url: {_esc(op.get('url'))}</div>")
        if op.get("error"):
            out.append(f"<div class='err'>{_esc(op.get('error'))}</div>")
        if op.get("request"):
            out.append(f"<div class='kv'>请求体:</div><pre>{_esc(_json_pretty(op.get('request')))}</pre>")
        if op.get("response"):
            out.append(f"<div class='kv'>响应体:</div><pre>{_esc(_json_pretty(op.get('response')))}</pre>")
    else:  # db_setup / db_checks
        action: str = op.get("action", "")
        out.append(f"<div class='ophead'>DB [{_esc(op.get('type'))}] "
                   f"<span class='s'>{_esc(op.get('profile'))}.{_esc(op.get('table'))} {_esc(action)}"
                   f" ｜ 返回 {_esc(op.get('rows_count'))} 行</span></div>")
        extra: list[str] = []
        if op.get("where"): extra.append(f"where={_json_pretty(op.get('where'))}")
        if op.get("inject"): extra.append(f"inject={_json_pretty(op.get('inject'))}")
        if op.get("expected") is not None: extra.append(f"expected={op.get('expected')}")
        if extra:
            out.append(f"<div class='kv'>{_esc(' ｜ '.join(extra))}</div>")
        if op.get("error"):
            out.append(f"<div class='err'>{_esc(op.get('error'))}</div>")
        preview: list | None = op.get("row_preview")
        if preview:
            out.append("<div class='rows'>行预览:</div>")
            cols: list[str] = list(preview[0].keys())
            out.append("<table><tr>" + "".join(f"<th>{_esc(c)}</th>" for c in cols) + "</tr>")
            for row in preview:
                out.append("<tr>" + "".join(f"<td>{_esc(_json_pretty(row.get(c)))}</td>" for c in cols) + "</tr>")
            out.append("</table>")
    out.append("</div>")
    return "".join(out)

def _op_list_item(op: dict, step_name: str, op_id: str) -> tuple[str, str, str]:
    """下拉面板条目: (短标签, title 全文, 跳转锚点 id)。短标签单行省略号截断, 全文悬停查看。"""
    if op.get("type") == "http":
        label: str = f"HTTP {op.get('method')} {op.get('path')} → {op.get('status_code')}"
        full: str = (f"步骤: {step_name}\n"
                     f"HTTP {op.get('method')} {op.get('path')} → {op.get('status_code')}｜耗时 {op.get('elapsed')}s\n"
                     f"url: {op.get('url') or ''}")
        if op.get("error"):
            full += f"\n错误: {op['error']}"
    else:
        action: str = (op.get("action") or "").upper()
        label = f"DB [{op.get('type')}] {op.get('table')} {action}"
        full = (f"步骤: {step_name}\n"
                f"DB [{op.get('type')}] {op.get('profile')}.{op.get('table')} {action}｜返回 {op.get('rows_count')} 行")
        if op.get("error"):
            full += f"\n错误: {op['error']}"
    return label, full, op_id

def _drop_panel_html(kind: str, items: list[tuple], title: str) -> str:
    """渲染一个隐藏的下拉面板(通过/失败)。默认 display:none, JS 切换显示。"""
    if items:
        body: str = "".join(
            f"<div class='drop-item' data-target='{_esc(t)}' title='{_esc(full)}'>{_esc(label)}</div>"
            for label, full, t in items
        )
    else:
        body = "<div class='drop-empty'>暂无内容</div>"
    return (f"<div class='drop' id='drop-{kind}' style='display:none'>"
            f"<div class='drop-head'><span>{_esc(title)} <span class='cnt'>({len(items)})</span></span>"
            f"<button class='drop-close' type='button' aria-label='关闭'>×</button></div>"
            f"<div class='drop-list'>{body}</div></div>")

# ═══════════════════════════════════════════════════════════════
#  Excel 报告 (.xlsx, openpyxl)
# ═══════════════════════════════════════════════════════════════
def generate_excel(result: StandardFlowResult, path: str) -> None:
    """生成 3-Sheet xlsx: 概览 / 步骤明细 / 操作明细。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    header_fill = PatternFill("solid", fgColor="F3F4F6")
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    def _style_header(ws) -> None:
        for cell in ws[1]:
            cell.font = bold
            cell.fill = header_fill

    def _fill_status(ws, row: int, passed: bool) -> None:
        cell = ws.cell(row=row, column=1)
        cell.fill = pass_fill if passed else fail_fill

    # ── 概览 ──
    ws = wb.active
    ws.title = "概览"
    ws.append(["项", "值"])
    ws.append(["流程名称", result.flow_name])
    ws.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws.append(["步骤数", result.total])
    ws.append(["通过", result.passed])
    ws.append(["失败", result.failed])
    ws.append(["通过率", f"{round(result.passed / result.total * 100, 1) if result.total else 0}%"])
    ws.append(["耗时(秒)", round(result.duration, 3)])
    ws.append(["是否通过", "是" if result.is_passed else "否"])
    _style_header(ws)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 60

    # ── 步骤明细 ──
    ws2 = wb.create_sheet("步骤明细")
    ws2.append(["流程", "步骤", "状态", "提取字段", "错误", "操作类型"])
    for step in result.steps:
        ws2.append([
            result.flow_name,
            step.name,
            "PASS" if step.passed else "FAIL",
            _json_pretty(step.extracted_vars, _EXCEL_MAX_CELL) if step.extracted_vars else "",
            "；".join(step.errors) if step.errors else "",
            "，".join(op.get("type", "?") for op in step.operations),
        ])
        _fill_status(ws2, ws2.max_row, step.passed)
    _style_header(ws2)
    for col, w in zip("ABCDEF", (18, 28, 10, 60, 60, 24)):
        ws2.column_dimensions[col].width = w
    for row in ws2.iter_rows(min_row=2):
        for c in row: c.alignment = wrap

    # ── 操作明细 ──
    ws3 = wb.create_sheet("操作明细")
    ws3.append(["流程", "步骤", "操作类型", "详情", "请求体", "响应体", "返回行数", "通过", "错误"])
    for step in result.steps:
        for op in step.operations:
            ws3.append([
                result.flow_name,
                step.name,
                op.get("type", "?"),
                _excel_op_detail(op),
                _json_pretty(op.get("request"), _EXCEL_MAX_CELL) if op.get("request") else "",
                _json_pretty(op.get("response"), _EXCEL_MAX_CELL) if op.get("response") else "",
                op.get("rows_count", ""),
                "是" if op.get("passed") else "否",
                op.get("error") or "",
            ])
            ws3.cell(row=ws3.max_row, column=8).fill = pass_fill if op.get("passed") else fail_fill
    _style_header(ws3)
    for col, w in zip("ABCDEFGHI", (18, 28, 12, 60, 90, 90, 12, 8, 60)):
        ws3.column_dimensions[col].width = w
    for row in ws3.iter_rows(min_row=2):
        for c in row: c.alignment = wrap

    wb.save(path)

def _excel_op_detail(op: dict) -> str:
    """操作摘要: HTTP 的 request/response 有独立列展示, 此处不再重复携带。"""
    d: dict = dict(op)
    if d.get("type") == "http":
        d.pop("request", None)
        d.pop("response", None)
    return _json_pretty(d, _EXCEL_MAX_CELL)
